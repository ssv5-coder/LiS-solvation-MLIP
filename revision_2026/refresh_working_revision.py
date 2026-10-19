#!/usr/bin/env python3
"""Apply source-audit corrections only to the generated working revision."""
import json
import sys
from common import WORK, save_json


def main():
    folder = WORK / "manuscript_working"
    marker = folder / "source_audit_labels_corrected.json"
    if marker.exists():
        raise FileExistsError("Working labels already corrected; review later changes manually")
    changes = []
    for filename in ["main.tex", "SI.tex"]:
        path = folder / filename
        text = path.read_text()
        for before, after in [
            ("two fully held-out AIMD trajectories", "two AIMD trajectories in the historical Testing archive, which overlaps training-time validation"),
            ("fully held-out PS and non-PS trajectories", "PS and non-PS historical benchmark trajectories"),
            ("held-out AIMD test trajectories", "historical AIMD benchmark trajectories"),
            ("held-out PS and non-PS test sets", "historical PS and non-PS benchmark sets"),
            ("held-out test sets", "historical benchmark sets"),
            ("held-out test set", "historical benchmark set"),
            ("held-out trajectories", "historical benchmark trajectories"),
            ("held-out trajectory errors", "historical benchmark trajectory errors"),
            ("held-out and transferability", "historical and transferability"),
            ("configurations not encountered during fine-tuning", "configurations in the historical benchmark archive, noting its overlap with training-time validation"),
        ]:
            if before in text:
                changes.append(dict(file=filename, before=before, after=after, occurrences=text.count(before)))
                text = text.replace(before, after)
        if filename == "main.tex":
            paragraphs = text.splitlines()
            old = next(p for p in paragraphs if p.startswith("Across the combined historical benchmark sets"))
            new = r"The archived training-time validation logs report stress MAE/RMSE values of 0.330/0.454 meV~\AA{}$^{-3}$ (Table~\ref{tab:validation_metrics}); these are not independent held-out metrics. The corresponding historical energy metric requires reconciliation because the preprocessing notebook and installed loss appear to normalize the target twice. We therefore do not interpret the logged 6.82/9.47 values as verified physical energy MAE/RMSE in the revised benchmark. The reported transferability energy MAE/RMSE of 5.56/8.06 meV atom$^{-1}$ describes composition-offset-corrected energies, not absolute total-energy accuracy. Raw absolute and shared-reference relative errors must be reported separately. Force diagnostics remain available in Figure~\ref{fig:SI_parity}, but numerical agreement must not be equated with independent validation of reactive kinetics."
            changes.append(dict(file=filename, before=old, after=new, occurrences=1))
            text = text.replace(old, new)
        else:
            note = r"\textcolor{blue}{Source-audit warning: inherited energy tables and figures in this working SI are retained for traceability, not accepted as corrected physical-energy benchmarks. Historical Testing trajectories overlap validation; raw total-energy normalization remains unresolved. Do not submit until these tables and their labels are reconciled with the revision audit.}" + "\n"
            text = text.replace(r"\begin{document}", r"\begin{document}"+"\n"+note, 1)
        path.write_text(text)
    save_json(marker, changes)
    sys.path.insert(0, str(WORK / "document_tools"))
    from docx import Document
    path = WORK / "Response_to_Reviewers_Working_20261019.docx"
    doc = Document(path)
    for p in doc.paragraphs:
        if "a bounded retry is being prepared" in p.text:
            p.text = p.text.replace("a bounded retry is being prepared", "a two-node retry also stopped in MPI before convergence. A final bounded one-node resource pilot uses increased memory and NCORE=30 without changing the physical model")
    doc.add_heading("Execution Update", 1)
    doc.add_paragraph("The complete original/optimized evaluator test passed on A100: 4.93-fold median inference speedup (1.340 to 0.272 s), with energy/force/stress parity and a 20-step trajectory difference below 0.000001 Angstrom. This is a short benchmark, not a guarantee of full-run wall time. The original potential remains unchanged.")
    doc.add_paragraph("All 21 original-model cluster predictions were freshly recomputed on GPU and agree with the historical cache within numerical tolerance. The 400 newer transferability geometries are distinct and excluded from the recovered final-stage training geometry sizes; this does not retroactively make the historical Testing archive independent or resolve the energy-label discrepancy.")
    doc.add_paragraph("Both small-cell density pilots completed 20 electronically converged ionic steps. Bounded 2 ps DFT extensions and matched-composition ML controls have been submitted. Equilibrium-density claims await stationarity and uncertainty checks. Four replicas are submitted; actual progress is recorded in status.json. HSE and Filippo's FES remain incomplete.")
    doc.save(path)


if __name__ == "__main__":
    main()
