#!/usr/bin/env python3
"""Create a working manuscript copy and an evidence/status-based response draft."""
import csv
import json
from pathlib import Path
import shutil
import sys
import zipfile
from common import BASE, WORK, REPO, save_json

sys.path.insert(0, str(WORK / "document_tools"))
from docx import Document
from docx.shared import Inches, Pt
from pypdf import PdfReader


def main():
    archive = next((BASE / "JPCRevision/editor").glob("*.zip"))
    folder = WORK / "manuscript_working"
    folder.mkdir(exist_ok=False)
    with zipfile.ZipFile(archive) as z:
        top = next(n.rsplit("/",1)[0] for n in z.namelist() if n.endswith("/main.tex") and not n.startswith("__MACOSX"))
        for name in z.namelist():
            if not name.startswith(top+"/") or name.endswith("/"):
                continue
            relative = Path(name[len(top)+1:])
            if relative.is_absolute() or ".." in relative.parts:
                raise ValueError("Unsafe archive member")
            target = folder / relative
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(z.read(name))
    original = (folder / "main.tex").read_text()
    (folder / "main_submitted_unchanged.tex").write_text(original)
    changes = []

    def replace(old, new, reason, expected=1):
        nonlocal original
        if original.count(old) != expected:
            raise ValueError(f"Expected {expected} matches for {old[:70]!r}; found {original.count(old)}")
        original = original.replace(old, new)
        changes.append(dict(reason=reason, old=old, new=new))

    paragraphs = original.split("\n")
    training = next(p for p in paragraphs if p.startswith("The pre-trained M3GNet model"))
    replace(training, r"The M3GNet potential was fine-tuned with the MatGL framework \cite{M3GNet,MGLDataset} by optimizing the complete network; no network layers were frozen. The archived final stage starts from an already adapted potential and contains 288,157 trainable network parameters. Its verified settings are 20 epochs, batch size 28 with accumulation over two batches, Adam with initial learning rate $10^{-4}$, the MatGL cosine scheduler with $T_{\max}=50$ and a minimum learning rate of $10^{-5}$, energy and force weights of 1.0, and a stress weight increasing from 0.01 to 0.1 over the 20 epochs. The neighbor and three-body cutoffs are 5.0 and 4.0~\AA{}, respectively. These settings describe the recovered final stage rather than the complete earlier training history. Original scripts, initialization, logs, and an executable reconstruction are available in the public repository." + "\n" + r"\textcolor{blue}{Revision audit: the historical validation selection overlaps the published AIMD Testing archive. These trajectories must not be described as fully independent until the split provenance is reconciled. Energy-label normalization is also under audit; the missing historical graph caches prevent exact reconstruction. The original production model has not been changed.}", "R2.1: source-backed final-stage settings and explicit unresolved split/normalization audit")
    replace("The observed mechanism is therefore stepwise,", "The temporal sequence in this trajectory is stepwise,", "R1.1: distinguish an observed sequence from a general mechanism")
    old = "Calculated DFT single-point energies along the MLIP trajectory further show that the sampled final state lies 0.40 eV below the lowest-energy intact configuration (Figure \\ref{fig:dynamics}B), indicating that formation of the separated products is energetically favorable. Accessing this state, however, involves overcoming a higher energy barrier by restructuring the local solvation environment, Li--S separation, charge redistribution, and eventual S--S cleavage. Finite-temperature fluctuations sampled during ML-MD provide access to these collective configurations and enable its spontaneous transformation. The computed energies should therefore be interpreted as sampled energy rises rather than quantitative activation barriers. Nonetheless, the DFT energetics provide an independent first-principles assessment of the configurations accessed by ML-MD, supporting the energetic plausibility of both the observed dissociation sequence and the resulting fragmented state."
    new = r"The static DFT energies of the selected, independently MLIP-relaxed cluster configurations place the sampled final configuration 0.40 eV below configuration B (Figure~\ref{fig:dynamics}). This snapshot difference is not a reaction free energy, and the intervening sampled energy rises are not activation barriers. Solvent configuration, cluster extraction, and separate geometry optimizations can all contribute to the profile. The event occurred during the pre-production heating segment, near a thermostat target of 338 K, rather than during stationary 300 K production. Consequently, its occurrence time cannot be interpreted as a measured 300 K rate constant or used to demonstrate thermal crossing of the plotted static energy rises. Within-state fluctuations, independent replicas, and a converged free-energy calculation are required before drawing kinetic or thermodynamic conclusions."
    replace(old, new, "R1.2/R1.4/R2.3: remove barrier and thermodynamic claims not established by snapshot energies")
    replace("Our results bring to light a sequential dissociation mechanism", "The original trajectory exhibits a candidate sequential dissociation pathway", "R1.1: qualify abstract scope")
    replace("which spin-polarized DFT identifies as an S$_3^{\\bullet-}$ radical pair", "whose spin-polarized PBE electronic structure is consistent with an S$_3^{\\bullet-}$ radical-pair assignment", "R1.5: distinguish PBE evidence from completed hybrid verification")
    # Correct unmistakable source typos; the exact journal line numbering is unavailable.
    for old, new in [("dynamicc oordination", "dynamic coordination"), ("etherelectrolytes", "ether electrolytes"), ("identities evolves", "identities evolve")]:
        if old in original:
            replace(old, new, "R2.4: source typo correction", original.count(old))
    marker = r"\section{Conclusions}"
    if marker not in original:
        marker = r"\section{Conclusion}"
    notes = r"""
% Revision additions are provisional until the pending jobs and FES are assessed.
\subsection{Revision Validation and Limitations}
The historical AIMD training trajectory contains 120,000 frames: 60,000 from 146-atom non-PS cells and 60,000 from 154-atom PS cells. All sulfur-pair distances were exported with periodic minimum-image treatment and frame provenance. Original polysulfide and TFSI sulfur membership is retained when interpreting these distances; geometric separation alone does not assign bond order or spin state.

Original-model predictions on 21 unrelaxed, fixed-composition 148-atom extracted clusters were compared with converged DFT single points at identical geometries. The current audit gives a force-component MAE of 185.84 meV~\AA{}$^{-1}$ and RMSE of 382.95 meV~\AA{}$^{-1}$. The raw absolute energy MAE is 4.6466 eV atom$^{-1}$, while the relative-energy MAE is 1.3247 eV per cluster when both methods are referenced to the first cluster. The large absolute offset and reactive-cluster errors require explicit treatment; these results must not be presented as confirmation of the bulk reaction barrier. Path-adapted-model results are in-sample and are reported separately from the original production model.

\textcolor{blue}{Pending: four independent replicas, hybrid-functional convergence and fragment-spin analysis, equilibrated density estimates, and Filippo's free-energy/within-state fluctuation analysis. No committee uncertainty is claimed while training-label provenance remains unresolved. The final scope of the mechanistic claims will depend on these results.}
"""
    if marker in original:
        original = original.replace(marker, notes+"\n"+marker, 1)
    else:
        original = original.replace(r"\end{document}", notes+"\n"+r"\end{document}")
    original = "% WORKING REVISION: NOT SUBMISSION-READY. Resolve blue audit notes and pending results.\n" + original
    (folder / "main.tex").write_text(original)
    save_json(folder / "wording_changes.json", changes)
    si = (folder / "SI.tex").read_text()
    (folder / "SI_submitted_unchanged.tex").write_text(si)
    for source, name in [(WORK/"training_coverage/training_ss_coverage.png", "revision_training_ss.png"),
                         (WORK/"reactive_validation/reactive_error_audit.png", "revision_reactive_errors.png")]:
        shutil.copy2(source, folder/"Pics"/name)
    appendix = r"""
\section{Revision Audit: Training Coverage and Reactive Errors}
The full sulfur-distance export covers 120,000 training frames. The starting sulfur membership is assigned once per source block using the TFSI S--O environment and retained across that block. Distances are minimum-image distances, not bond orders.
\begin{figure}[htbp]\centering
\includegraphics[width=\textwidth]{Pics/revision_training_ss.png}
\caption{Training sulfur-pair distance distributions. The right panel follows pairs initially within 2.6~\AA{}; this includes incidental close contacts and is not a chemical bond assignment.}
\label{fig:revision_training_ss}\end{figure}
\begin{figure}[htbp]\centering
\includegraphics[width=\textwidth]{Pics/revision_reactive_errors.png}
\caption{Original production-model errors against converged DFT for identical unrelaxed 148-atom extracted clusters. Relative energies use the same first-frame reference. Breaking-pair force errors concern original periodic-system atom IDs 3432 and 3433. These cluster tests do not establish a bulk activation free energy.}
\label{fig:revision_reactive_errors}\end{figure}
\paragraph{Density comparison.} The input audit located 88 historical AIMD stage inputs and no constant-300 K NPT stage among them. NVT densities are imposed by the cell; temperature-ramp endpoints are not equilibrated 300 K density estimates. New small-cell pilots therefore separate timing/stability checks from subsequent equilibrium sampling. Natural-mass and simulation-mass densities are retained separately.
\paragraph{Pending results.} Replica statistics, HSE spin comparisons, equilibrated DFT/ML density estimates, and converged FES results are not yet supplied by this working supplement.
"""
    (folder / "SI.tex").write_text(si.replace(r"\end{document}", appendix+"\n"+r"\end{document}"))
    responses = [
        ("R1.1", "Independent trajectories and event frequency", "Four independently packed 3,445-atom PS replicas have been prepared with distinct packing, velocity, and thermostat seeds. The nominal original annealing/production protocol is retained. Both initial polysulfides will be analyzed using predeclared geometric thresholds, persistence, sensitivity checks, and right censoring. Replica jobs are pending/running, so no event-frequency result is claimed. The working text describes the original sequence as a candidate pathway."),
        ("R1.2", "Barrier versus observed timescale", "We thank the reviewer for identifying this inconsistency. The source audit places cleavage during heating near 338 K, not stationary 300 K MD. Static energies of extracted and independently relaxed clusters are not activation free energies. The working text removes the thermal-barrier-crossing argument. Filippo's FES remains pending; comparison to Han and Xu must match reaction definition, temperature, solvent and ensemble."),
        ("R1.3", "Training coverage and reactive accuracy", "The training cells contain 146/154 atoms, not 30-40. All 120,000 training frames now have a complete sulfur-pair distance export and coverage figure. A raw-DFT audit of 21 unrelaxed clusters gives original-model force MAE/RMSE of 0.18584/0.38295 eV/Angstrom. Per-frame, per-element and breaking-pair errors are included. Large absolute energy offsets and label normalization remain under audit. Full-system 3,445-atom AIMD is outside the revision scope, and cluster validation is not presented as a replacement for bulk AIMD. Committee training is held because the prerequisite normalization audit has not passed."),
        ("R1.4", "Snapshot fluctuations and thermodynamic favorability", "Filippo leads within-state energy fluctuation analysis and FES. The working manuscript no longer interprets the 0.40 eV difference between selected configurations as a thermodynamic reaction energy or free energy. Final response must insert his converged results or retain the limitation."),
        ("R1.5", "Spin treatment and hybrid confirmation", "Training AIMD inputs are spin-unpolarized; final-state PBE spin analysis is a separate calculation. A static HSE06 pilot retains the 148-atom final geometry and original vacuum cell with 600 eV cutoff. AFM, FM, and nonmagnetic input variants are prepared. HSE06-specific D3 zero-damping parameters are explicit. The initial HSE pilot exceeded a one-node memory allocation; a bounded retry is being prepared. No new hybrid radical confirmation is claimed before convergence and spin-density analysis."),
        ("R2.1", "Fine-tuning explanation and code", "We have made the training script, final-stage initialization, historical script/logs, full-network trainability, source/build provenance and selected genuine-velocity checkpoints available in the GitHub reproducibility release. The working Methods now distinguish the verified 20-epoch final stage from prior adaptation. Dataset overlap and energy-label reconstruction limitations are explicitly disclosed."),
        ("R2.2", "Density overprediction and slow diffusion", "An 88-input AIMD ensemble audit found no constant-300 K NPT stage. We therefore distinguish imposed NVT volumes from equilibrium DFT predictions and retain natural and simulation masses separately. Original ML-MD changes hydrogen mass between annealing and final production. Bounded new 300 K NPT pilots and matched-composition ML controls are prepared/submitted. More representative training data may improve sampling-domain errors, but cannot guarantee correction of the DFT reference or establish overbinding as the sole cause."),
        ("R2.3", "Path accuracy and physical meaning", "The revision combines original-model per-frame errors, replica statistics, spin tests and Filippo's FES, keeping their evidential roles separate. Current reactive-cluster errors do not independently establish physical bulk kinetics. The path-adapted fit is not used to validate the original production event. Final conclusions remain conditional on pending tests."),
        ("R2.4", "Page 3, line 31 typo", "The archived local PDF has no journal line numbering. The working copy corrects obvious source spacing/grammar errors (including 'dynamicc oordination' and 'etherelectrolytes' where present). The exact line-31 issue must be checked against the journal-generated proof before marking this response complete."),
    ]
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("JPC Revision: Evidence and Response Draft", 0)
    doc.add_paragraph("19 October 2026. Working document, not a completed rebuttal. Pending calculations are explicitly labeled; page/line references must be updated after the revised manuscript is finalized.")
    doc.add_paragraph("We thank both reviewers for their careful assessment and constructive suggestions. We have organized the revision around reproducibility, the scope of the mechanistic evidence, and direct tests of the reactive configurations. The responses below separate completed audits from ongoing work.")
    for code, title, response in responses:
        doc.add_heading(code+": "+title, 1)
        doc.add_paragraph(response)
        doc.add_paragraph("Location: manuscript_working/main.tex and SI.tex; final page/line numbering pending.")
    doc.add_heading("Critical Audit Findings", 1)
    diagnostic = WORK / "training_audit/normalization_summary.json"
    if diagnostic.exists():
        values = json.loads(diagnostic.read_text())
        for row in values["results"]:
            doc.add_paragraph(f"40-frame {row['subset']} diagnostic: standard total-energy MAE {row['standard_energy_mae_meV_atom']:.2f} meV/atom; historical double-normalized diagnostic {row['double_normalized_energy_mae']:.2f}. This is not a full benchmark or proof of the missing-cache contents. Do not retain the original energy-accuracy claim without resolving the discrepancy.")
    doc.add_paragraph("The original model, submitted manuscript, and submitted SI are preserved. No existing scientific result has been silently replaced by a corrected-model calculation.")
    doc.save(WORK / "Response_to_Reviewers_Working_20261019.docx")
    save_json(WORK / "reviewer_response_tracker.json", [dict(comment=c, topic=t, response=r, status="working / evidence and pending tasks distinguished") for c,t,r in responses])
    # Retain page 3 text to locate the reviewer typo without guessing journal line numbers.
    pdf = BASE / "JPCRevision/editor/MLIP_LiS_Disovery-3.pdf"
    (WORK / "submitted_page3_text.txt").write_text(PdfReader(pdf).pages[2].extract_text())


if __name__ == "__main__":
    main()
